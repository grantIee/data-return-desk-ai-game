"""Bot player for The Return Desk. Analyzes files and applies rules."""
import sys
import json
import time
import random
import requests
from pathlib import Path
from io import BytesIO

BASE = "http://localhost:8888"


def analyze_pdf(content: bytes) -> bool:
    """Check if PDF contains a customer photo. Look for image markers."""
    # fpdf2 embeds images as /Subtype /Image in the PDF stream
    return b"/Subtype /Image" in content


def analyze_xlsx(content: bytes) -> dict:
    """Parse transaction history and calculate return rate + dollar ratio."""
    from openpyxl import load_workbook
    wb = load_workbook(BytesIO(content), read_only=True)
    ws = wb.active

    purchases = 0
    returns = 0
    purchase_amount = 0.0
    return_amount = 0.0

    for row in ws.iter_rows(min_row=4, values_only=True):
        if not row or not row[3]:
            continue
        txn_type = str(row[3]).strip()
        amount = float(row[2]) if row[2] else 0

        if txn_type == "Purchase":
            purchases += 1
            purchase_amount += amount
        elif txn_type == "Return":
            returns += 1
            return_amount += amount

    total = purchases + returns
    return_rate = returns / total if total > 0 else 0
    dollar_ratio = return_amount / purchase_amount if purchase_amount > 0 else 0

    return {
        "purchases": purchases,
        "returns": returns,
        "return_rate": return_rate,
        "purchase_amount": purchase_amount,
        "return_amount": return_amount,
        "dollar_ratio": dollar_ratio,
    }


def analyze_pptx(content: bytes) -> dict:
    """Find the assessment slide in the fraud report."""
    from pptx import Presentation
    prs = Presentation(BytesIO(content))

    fraud_assessment = "unknown"
    spending_potential = "unknown"

    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                text = shape.text_frame.text.lower()
                if "assessment:" in text:
                    if "not malicious" in text:
                        fraud_assessment = "not_malicious"
                    elif "suspicious" in text:
                        fraud_assessment = "suspicious"
                    elif "flagged" in text:
                        fraud_assessment = "flagged"
                if "spending potential:" in text:
                    if "high" in text:
                        spending_potential = "high"
                    elif "medium" in text:
                        spending_potential = "medium"
                    elif "low" in text:
                        spending_potential = "low"

    return {"fraud_assessment": fraud_assessment, "spending_potential": spending_potential}


def decide(has_photo: bool, fraud: str, spending: str, return_rate: float,
           return_amount: float, purchase_amount: float, loyalty_tier: str) -> str:
    """Apply the decision rules."""
    # Rule 1: Photo Gate
    if not has_photo:
        return "DENY"
    # Rule 2: Loyalty Override
    if loyalty_tier == "GOLD":
        return "ACCEPT"
    # Rule 3: Fraud Override
    if fraud == "not_malicious" and spending == "high":
        return "ACCEPT"
    # Rule 4 + 5: Return Rate with Dollar Exception
    if return_rate > 0.20:
        if return_amount < 0.80 * purchase_amount:
            return "ACCEPT"
        return "DENY"
    # Rule 5 standalone
    if return_amount < 0.80 * purchase_amount:
        return "ACCEPT"
    # Rule 6: Default
    return "DENY"


def play(session_id: str, num_customers: int = 10, skill_level: str = "smart"):
    """Play the game for a session."""
    # Load BQ lookup for loyalty tier (smart bots have this)
    bq_lookup = {}
    bq_file = Path(__file__).parent / "bq_load.jsonl"
    if bq_file.exists() and skill_level in ("smart", "agentic"):
        with open(bq_file) as f:
            for line in f:
                row = json.loads(line)
                bq_lookup[row["customer_id"]] = row.get("loyalty_tier", "NONE")

    i = 0
    while True:
        # Get next customer (stop if game ended or error)
        resp = requests.get(f"{BASE}/api/session/{session_id}/next")
        if resp.status_code != 200:
            break
        i += 1
        data = resp.json()
        cid = data["customer_id"]

        if skill_level == "random":
            # Random bot — just guesses
            decision = random.choice(["ACCEPT", "DENY"])
            time.sleep(random.uniform(0.5, 2.0))
        elif skill_level == "naive":
            # Reads files but doesn't know about BQ
            pdf = requests.get(f"{BASE}/api/files/{cid}/receipt.pdf").content
            xlsx = requests.get(f"{BASE}/api/files/{cid}/transactions.xlsx").content
            pptx = requests.get(f"{BASE}/api/files/{cid}/fraud_report.pptx").content
            has_photo = analyze_pdf(pdf)
            tx = analyze_xlsx(xlsx)
            fraud = analyze_pptx(pptx)
            decision = decide(has_photo, fraud["fraud_assessment"], fraud["spending_potential"],
                              tx["return_rate"], tx["return_amount"], tx["purchase_amount"],
                              "NONE")  # doesn't know about BQ
            time.sleep(random.uniform(1.0, 3.0))
        elif skill_level == "smart":
            # Reads files AND has BQ data
            pdf = requests.get(f"{BASE}/api/files/{cid}/receipt.pdf").content
            xlsx = requests.get(f"{BASE}/api/files/{cid}/transactions.xlsx").content
            pptx = requests.get(f"{BASE}/api/files/{cid}/fraud_report.pptx").content
            has_photo = analyze_pdf(pdf)
            tx = analyze_xlsx(xlsx)
            fraud = analyze_pptx(pptx)
            loyalty = bq_lookup.get(cid, "NONE")
            decision = decide(has_photo, fraud["fraud_assessment"], fraud["spending_potential"],
                              tx["return_rate"], tx["return_amount"], tx["purchase_amount"],
                              loyalty)
            time.sleep(random.uniform(0.3, 1.5))
        elif skill_level == "agentic":
            # Fast + accurate — full analysis, minimal delay
            pdf = requests.get(f"{BASE}/api/files/{cid}/receipt.pdf").content
            xlsx = requests.get(f"{BASE}/api/files/{cid}/transactions.xlsx").content
            pptx = requests.get(f"{BASE}/api/files/{cid}/fraud_report.pptx").content
            has_photo = analyze_pdf(pdf)
            tx = analyze_xlsx(xlsx)
            fraud = analyze_pptx(pptx)
            loyalty = bq_lookup.get(cid, "NONE")
            decision = decide(has_photo, fraud["fraud_assessment"], fraud["spending_potential"],
                              tx["return_rate"], tx["return_amount"], tx["purchase_amount"],
                              loyalty)
            time.sleep(random.uniform(0.1, 0.5))

        # Submit
        resp = requests.post(f"{BASE}/api/session/{session_id}/decide",
                             json={"customer_id": cid, "decision": decision})
        if resp.status_code != 200:
            break
        result = resp.json()

    # Print final stats
    resp = requests.get(f"{BASE}/api/session/{session_id}")
    if resp.status_code == 200:
        stats = resp.json()
        print(f"{stats['name']}: {stats['correct']}/{stats['total']} correct, "
              f"accuracy={stats['accuracy']}%, avg={stats['avg_time']}s, score={stats['score']}")


if __name__ == "__main__":
    session_id = sys.argv[1]
    num = int(sys.argv[2]) if len(sys.argv) > 2 else 10
    skill = sys.argv[3] if len(sys.argv) > 3 else "smart"
    play(session_id, num, skill)
